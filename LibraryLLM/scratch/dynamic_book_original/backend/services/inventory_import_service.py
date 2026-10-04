import csv
import io
import os
import uuid
from datetime import datetime

import duckdb
from openpyxl import load_workbook

from backend.database.mongodb import (
    books_collection,
    library_inventory_collection,
    inventory_preview_collection,
)


ISBN_DB = r"D:\SDC\LibraryLLM\datasets\ai\mappings\isbn_lookup.duckdb"


REQUIRED_COLUMNS = {
    "isbn",
    "total_copies",
    "available_copies"
}

OPTIONAL_COLUMNS = {
    "title",
    "author",
    "shelf_location"
}


def normalize_isbn(value):

    if value is None:
        return None

    value = str(value).strip()

    if not value:
        return None

    normalized = "".join(
        character
        for character in value
        if character.isdigit() or character.upper() == "X"
    )

    return normalized or None


def read_csv_file(file_bytes):

    text = file_bytes.decode("utf-8-sig")

    reader = csv.DictReader(
        io.StringIO(text)
    )

    if not reader.fieldnames:
        raise ValueError(
            "CSV file has no header row"
        )

    return reader.fieldnames, list(reader)


def read_xlsx_file(file_bytes):

    workbook = load_workbook(
        filename=io.BytesIO(file_bytes),
        read_only=True,
        data_only=True
    )

    worksheet = workbook.active

    rows = list(
        worksheet.iter_rows(
            values_only=True
        )
    )

    workbook.close()

    if not rows:
        raise ValueError(
            "Excel file is empty"
        )

    headers = [
        str(value).strip()
        if value is not None
        else ""
        for value in rows[0]
    ]

    records = []

    for row in rows[1:]:

        record = {}

        for index, header in enumerate(headers):

            if not header:
                continue

            record[header] = (
                row[index]
                if index < len(row)
                else None
            )

        records.append(record)

    return headers, records


def read_uploaded_file(
    filename,
    file_bytes
):

    extension = os.path.splitext(
        filename
    )[1].lower()

    if extension == ".csv":

        return read_csv_file(
            file_bytes
        )

    if extension == ".xlsx":

        return read_xlsx_file(
            file_bytes
        )

    raise ValueError(
        "Only CSV and XLSX files are supported"
    )


def get_isbn_matches(isbns):

    if not isbns:
        return {}

    connection = duckdb.connect(
        ISBN_DB,
        read_only=True
    )

    placeholders = ",".join(
        ["?"] * len(isbns)
    )

    query = f"""
        SELECT
            isbn,
            work_id,
            title,
            authors
        FROM isbn_lookup
        WHERE isbn IN ({placeholders})
    """

    rows = connection.execute(
        query,
        list(isbns)
    ).fetchall()

    connection.close()

    matches = {}

    for isbn, work_id, title, authors in rows:

        matches.setdefault(
            isbn,
            []
        ).append(
            {
                "work_id": work_id,
                "title": title,
                "authors": authors
            }
        )

    return matches


def create_preview(
    library_id,
    filename,
    file_bytes
):

    headers, rows = read_uploaded_file(
        filename,
        file_bytes
    )

    normalized_headers = {
        str(header).strip().lower()
        for header in headers
        if header
    }

    missing_columns = (
        REQUIRED_COLUMNS
        - normalized_headers
    )

    if missing_columns:

        raise ValueError(
            "Missing required columns: "
            + ", ".join(
                sorted(missing_columns)
            )
        )

    prepared_rows = []

    isbn_set = set()

    for row_number, row in enumerate(
        rows,
        start=2
    ):

        normalized_row = {
            str(key).strip().lower(): value
            for key, value in row.items()
            if key is not None
        }

        isbn = normalize_isbn(
            normalized_row.get("isbn")
        )

        try:

            total_copies = int(
                normalized_row.get(
                    "total_copies",
                    1
                )
            )

            available_copies = int(
                normalized_row.get(
                    "available_copies",
                    total_copies
                )
            )

        except (
            TypeError,
            ValueError
        ):

            prepared_rows.append(
                {
                    "row_number": row_number,
                    "status": "INVALID",
                    "reason": "Invalid copy count",
                    "isbn": isbn
                }
            )

            continue

        if not isbn:

            prepared_rows.append(
                {
                    "row_number": row_number,
                    "status": "INVALID",
                    "reason": "Missing ISBN",
                    "isbn": None
                }
            )

            continue

        if total_copies < 1:

            prepared_rows.append(
                {
                    "row_number": row_number,
                    "status": "INVALID",
                    "reason": "total_copies must be at least 1",
                    "isbn": isbn
                }
            )

            continue

        if (
            available_copies < 0
            or available_copies > total_copies
        ):

            prepared_rows.append(
                {
                    "row_number": row_number,
                    "status": "INVALID",
                    "reason": "Invalid available_copies",
                    "isbn": isbn
                }
            )

            continue

        prepared_rows.append(
            {
                "row_number": row_number,
                "isbn": isbn,
                "title": normalized_row.get(
                    "title"
                ),
                "author": normalized_row.get(
                    "author"
                ),
                "total_copies": total_copies,
                "available_copies": available_copies,
                "shelf_location": normalized_row.get(
                    "shelf_location"
                )
            }
        )

        isbn_set.add(isbn)

    # ---------------------------------------------------------
    # ISBN matching
    # ---------------------------------------------------------

    matches = get_isbn_matches(
        isbn_set
    )

    matched_work_ids = set()

    for row in prepared_rows:

        if row.get("status") == "INVALID":
            continue

        isbn = row["isbn"]

        candidates = matches.get(
            isbn,
            []
        )

        if len(candidates) == 0:

            row["status"] = "NOT_FOUND"
            row["reason"] = (
                "ISBN not found in LuminaR catalogue"
            )

            continue

        if len(candidates) > 1:

            row["status"] = "AMBIGUOUS"
            row["reason"] = (
                "ISBN maps to multiple works"
            )

            row["candidates"] = candidates

            continue

        match = candidates[0]

        work_id = match["work_id"]

        # Verify the work exists in MongoDB.
        book = books_collection.find_one(
            {"work_id": work_id},
            {
                "_id": 0,
                "work_id": 1,
                "title": 1,
                "authors": 1
            }
        )

        if book is None:

            row["status"] = "NOT_FOUND"
            row["reason"] = (
                "Mapped work does not exist "
                "in LuminaR catalogue"
            )

            continue

        existing = (
            library_inventory_collection.find_one(
                {
                    "library_id": library_id,
                    "work_id": work_id
                }
            )
        )

        row["work_id"] = work_id
        row["catalogue_title"] = (
            book.get("title")
        )
        row["catalogue_authors"] = (
            book.get("authors")
        )

        if existing:

            row["status"] = "EXISTS"
            row["reason"] = (
                "Book already exists in "
                "library inventory"
            )

        else:

            row["status"] = "MATCHED"
            row["reason"] = "Exact ISBN match"

            matched_work_ids.add(
                work_id
            )

    # ---------------------------------------------------------
    # Summary
    # ---------------------------------------------------------

    summary = {
        "total_rows": len(rows),
        "matched": sum(
            1
            for row in prepared_rows
            if row.get("status") == "MATCHED"
        ),
        "already_exists": sum(
            1
            for row in prepared_rows
            if row.get("status") == "EXISTS"
        ),
        "not_found": sum(
            1
            for row in prepared_rows
            if row.get("status") == "NOT_FOUND"
        ),
        "ambiguous": sum(
            1
            for row in prepared_rows
            if row.get("status") == "AMBIGUOUS"
        ),
        "invalid": sum(
            1
            for row in prepared_rows
            if row.get("status") == "INVALID"
        )
    }

    # ---------------------------------------------------------
    # Save preview to MongoDB
    # ---------------------------------------------------------

    preview_id = str(uuid.uuid4())

    preview_document = {
        "preview_id": preview_id,
        "library_id": library_id,
        "filename": filename,
        "summary": summary,
        "rows": prepared_rows,
        "status": "PENDING",
        "created_at": datetime.utcnow(),
    }

    inventory_preview_collection.insert_one(
        preview_document
    )

    return {
        "preview_id": preview_id,
        "library_id": library_id,
        "filename": filename,
        "summary": summary,
        "rows": prepared_rows
    }


# =========================================================
# CONFIRM INVENTORY IMPORT
# =========================================================

def confirm_inventory_import(
    preview_id,
    library_id
):
    """
    Confirm a previously generated inventory preview.

    Only MATCHED rows are imported.

    EXISTS, NOT_FOUND, AMBIGUOUS and INVALID
    rows are skipped.

    Multiple MATCHED rows pointing to the same
    work_id are consolidated into one inventory record.
    """

    preview = inventory_preview_collection.find_one(
        {
            "preview_id": preview_id,
            "library_id": library_id,
        },
        {
            "_id": 0
        }
    )

    if not preview:

        raise ValueError(
            "Preview not found or does not belong "
            "to this library"
        )

    # Prevent duplicate confirmation
    if preview.get("status") == "CONFIRMED":

        raise ValueError(
            "This preview has already been confirmed"
        )

    rows = preview.get(
        "rows",
        []
    )

    # ---------------------------------------------------------
    # Group MATCHED rows by work_id
    # ---------------------------------------------------------

    grouped = {}

    for row in rows:

        if row.get("status") != "MATCHED":
            continue

        work_id = row.get(
            "work_id"
        )

        if not work_id:
            continue

        if work_id not in grouped:

            grouped[work_id] = {
                "work_id": work_id,
                "isbn": row.get("isbn"),
                "title": row.get("title"),
                "author": row.get("author"),
                "total_copies": 0,
                "available_copies": 0,
                "shelf_location": row.get(
                    "shelf_location"
                ),
                "matched_rows": 0
            }

        grouped[work_id][
            "total_copies"
        ] += int(
            row.get(
                "total_copies"
            ) or 0
        )

        grouped[work_id][
            "available_copies"
        ] += int(
            row.get(
                "available_copies"
            ) or 0
        )

        grouped[work_id][
            "matched_rows"
        ] += 1

    # ---------------------------------------------------------
    # Counters
    # ---------------------------------------------------------

    imported = []
    consolidated = []

    skipped_existing = 0
    skipped_not_found = 0
    skipped_invalid = 0
    skipped_ambiguous = 0

    failed = []

    # ---------------------------------------------------------
    # Count skipped rows
    # ---------------------------------------------------------

    for row in rows:

        status = row.get(
            "status"
        )

        if status == "EXISTS":

            skipped_existing += 1

        elif status == "NOT_FOUND":

            skipped_not_found += 1

        elif status == "INVALID":

            skipped_invalid += 1

        elif status == "AMBIGUOUS":

            skipped_ambiguous += 1

    # ---------------------------------------------------------
    # Insert consolidated inventory
    # ---------------------------------------------------------

    for work_id, item in grouped.items():

        existing = (
            library_inventory_collection.find_one(
                {
                    "library_id": library_id,
                    "work_id": work_id,
                }
            )
        )

        if existing:

            skipped_existing += 1
            continue

        now = datetime.utcnow()

        document = {
            "library_id": library_id,
            "work_id": work_id,
            "isbn": item["isbn"],
            "total_copies": item[
                "total_copies"
            ],
            "available_copies": item[
                "available_copies"
            ],
            "shelf_location": item[
                "shelf_location"
            ],
            "source": "upload",
            "created_at": now,
            "updated_at": now,
        }

        try:

            library_inventory_collection.insert_one(
                document
            )

            imported.append(
                {
                    "work_id": work_id,
                    "isbn": item["isbn"],
                    "title": item["title"],
                    "total_copies": item[
                        "total_copies"
                    ],
                    "available_copies": item[
                        "available_copies"
                    ]
                }
            )

            # More than one CSV row mapped
            # to the same work.
            if item["matched_rows"] > 1:

                consolidated.append(
                    {
                        "work_id": work_id,
                        "rows_combined": item[
                            "matched_rows"
                        ]
                    }
                )

        except Exception as exc:

            failed.append(
                {
                    "work_id": work_id,
                    "error": str(exc)
                }
            )

    # ---------------------------------------------------------
    # Mark preview as confirmed
    # ---------------------------------------------------------

    inventory_preview_collection.update_one(
        {
            "preview_id": preview_id,
            "library_id": library_id,
        },
        {
            "$set": {
                "status": "CONFIRMED",
                "confirmed_at": datetime.utcnow()
            }
        }
    )

    # ---------------------------------------------------------
    # Return result
    # ---------------------------------------------------------

    return {
        "preview_id": preview_id,
        "library_id": library_id,
        "status": "CONFIRMED",

        "imported": len(
            imported
        ),

        "consolidated": len(
            consolidated
        ),

        "skipped_existing": skipped_existing,

        "skipped_not_found": skipped_not_found,

        "skipped_ambiguous": skipped_ambiguous,

        "skipped_invalid": skipped_invalid,

        "failed": len(
            failed
        ),

        "imported_books": imported,

        "consolidated_work_ids": consolidated,

        "errors": failed
    }