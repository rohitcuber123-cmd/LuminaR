"""Bounded CSV/XLSX previews bound to actor, UUID tenant, file hash and expiry."""
import csv
from datetime import timedelta, timezone
import hashlib
import io
import json
import re
from uuid import uuid4
import zipfile

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from backend.schemas.library_catalogue import BookCreate, CopyCreate
from backend.institution_dependencies import actor_id
from backend.services.library_inventory_access import require_library_staff
from backend.services.institution_repository import audit, library_mutation, now, require_readable
from backend.services.library_catalogue_repository import TenantInventory
from backend.services.library_catalogue_linking import isbn_candidates
from backend.services.library_catalogue_service import compatible, duplicate, insert_locked, normalized_title, prepare_book, add_copy_locked


MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 40
ALIASES = {
    'title': 'title', 'book_title': 'title', 'name': 'title',
    'subject': 'subjects', 'subjects': 'subjects', 'genre': 'subjects', 'category': 'subjects',
    'isbn': 'isbn', 'isbn10': 'isbn_10', 'isbn_10': 'isbn_10', 'isbn13': 'isbn_13', 'isbn_13': 'isbn_13',
    'work_id': 'global_work_id', 'global_work_id': 'global_work_id', 'author': 'authors', 'authors': 'authors',
    'publisher': 'publisher', 'year': 'publication_year', 'publication_year': 'publication_year',
    'edition': 'edition', 'language': 'language', 'description': 'description',
    'copies': 'copies', 'quantity': 'copies', 'total_copies': 'copies', 'available_copies': 'available_copies',
    'shelf': 'shelf_location', 'shelf_location': 'shelf_location', 'call_number': 'local_call_number', 'local_call_number': 'local_call_number',
    'tags': 'local_tags', 'local_tags': 'local_tags', 'acquisition_note': 'acquisition_note', 'barcode': 'barcode',
}
FORBIDDEN = {'library_id', 'canonical_library_id', 'library_book_id', 'holding_id', 'copy_id', 'accession_number',
             'accession_sequence', 'created_by', 'created_by_user_id', 'updated_by_user_id', 'role', 'status'}


def parse_file(filename, content):
    if len(content) > MAX_BYTES:
        raise HTTPException(413, 'Inventory files must be at most 5 MiB.')
    name = re.split(r'[/\\]', filename or 'inventory.csv')[-1][:120]
    extension = name.rsplit('.', 1)[-1].lower()
    workbook = None
    try:
        if extension == 'csv':
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                raise HTTPException(422, 'CSV must use UTF-8 encoding.')
            reader = csv.reader(io.StringIO(text), strict=True)
        elif extension == 'xlsx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(v.file_size for v in archive.infolist()) > 25 * 1024 * 1024:
                    raise HTTPException(413, 'Expanded workbook is too large.')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            reader = workbook.active.iter_rows(values_only=True)
        else:
            raise HTTPException(422, 'Only CSV and XLSX inventory files are supported.')
        raw_headers = next(reader, None)
        if not raw_headers or len(raw_headers) > MAX_COLUMNS:
            raise HTTPException(422, 'A header row with at most 40 columns is required.')
        headers = [str(v).strip().lower() if v is not None else '' for v in raw_headers]
        if len(set(headers)) != len(headers) or any(not h for h in headers):
            raise HTTPException(422, 'Column names must be nonempty and unique.')
        if set(headers) & FORBIDDEN:
            raise HTTPException(422, 'Identity, tenant, role and accession columns are not allowed in normal inventory imports.')
        keys = [ALIASES.get(h) for h in headers]
        recognized = [k for k in keys if k]
        if len(set(recognized)) != len(recognized):
            raise HTTPException(422, 'Multiple columns map to the same field; choose one explicit alias.')
        if not {'title', 'subjects'} <= set(recognized):
            raise HTTPException(422, 'Minimum columns: title and subject or genre.')
        rows = []
        for number, raw in enumerate(reader, 2):
            if number > MAX_ROWS + 1:
                raise HTTPException(413, 'Inventory imports support at most 2000 data rows.')
            if len(raw) > len(headers):
                raise HTTPException(422, f'Row {number} has more cells than the header.')
            if any(v is not None and len(str(v)) > 4000 for v in raw):
                raise HTTPException(422, f'Row {number} contains an oversized value.')
            row = {key: raw[i] for i, key in enumerate(keys) if key and i < len(raw) and raw[i] is not None and str(raw[i]).strip()}
            rows.append((number, row))
        if not rows:
            raise HTTPException(422, 'Inventory file has no data rows.')
        return name, rows
    except (csv.Error, zipfile.BadZipFile, ValueError, KeyError) as error:
        raise HTTPException(422, 'Unable to parse the inventory file. Check its format and columns.') from error
    finally:
        if workbook:
            workbook.close()


def validated_row(raw):
    if not str(raw.get('title') or '').strip():
        raise ValueError('MISSING_REQUIRED_TITLE')
    if not raw.get('subjects') or not str(raw['subjects']).strip():
        raise ValueError('MISSING_REQUIRED_SUBJECT')
    data = dict(raw)
    for field in ('copies', 'available_copies', 'publication_year'):
        if field in data:
            value = data[field]
            if isinstance(value, bool) or (isinstance(value, float) and not value.is_integer()) or not re.fullmatch(r'-?\d+', str(int(value) if isinstance(value, float) else value).strip()):
                raise ValueError(field + ' must be a whole number.')
            data[field] = int(value)
    # Numeric spreadsheet ISBNs are formatted as integer text, never guessed/padded.
    for field in ('isbn', 'isbn_10', 'isbn_13'):
        if field in data and isinstance(data[field], (int, float)) and not isinstance(data[field], bool):
            if isinstance(data[field], float) and not data[field].is_integer():
                raise ValueError('ISBN must not contain decimals.')
            data[field] = str(int(data[field]))
    return BookCreate.model_validate(data)


def safe_reason(error):
    if isinstance(error, ValidationError):
        first = error.errors()[0]
        return '.'.join(str(k) for k in first['loc']) + ': ' + first['msg'].removeprefix('Value error, ')
    if isinstance(error, HTTPException):
        return str(error.detail.get('message', 'Invalid row.')) if isinstance(error.detail, dict) else str(error.detail)
    return str(error)


def summary(rows):
    counts = {kind: sum(r['classification'] == kind for r in rows) for kind in ('NEW_LOCAL', 'GLOBAL_MATCH', 'EXISTING_HOLDING', 'POSSIBLE_DUPLICATE', 'INVALID')}
    return {'rows_total': len(rows), 'rows_valid': len(rows) - counts['INVALID'], 'rows_invalid': counts['INVALID'],
            'global_matches': counts['GLOBAL_MATCH'], 'new_library_records': counts['NEW_LOCAL'],
            'existing_holdings': counts['EXISTING_HOLDING'], 'possible_duplicates': counts['POSSIBLE_DUPLICATE'], 'classification_counts': counts}


def preview(library_id, filename, content, current_user):
    require_library_staff(library_id, current_user)
    name, raw_rows = parse_file(filename, content)
    prepared = []; numbers = []
    for number, raw in raw_rows:
        try:
            request = validated_row(raw)
            prepared.append({'row_number': number, 'request': request})
            data = request.model_dump(); numbers.extend(v for v in (data['isbn'], data['isbn_10'], data['isbn_13']) if v)
        except (ValueError, ValidationError) as error:
            prepared.append({'row_number': number, 'classification': 'INVALID', 'reason': safe_reason(error), 'title': str(raw.get('title') or '')})
    candidates = isbn_candidates(numbers)
    repo = TenantInventory(library_id); rows = []; pending = []
    for item in prepared:
        if 'request' not in item:
            rows.append(item); continue
        try:
            data = prepare_book(item['request'], candidates)
            match, classification = duplicate(repo, data)
            same = next((r for r in pending if (data.get('global_work_id') and data['global_work_id'] == r['data'].get('global_work_id')) or
                         (any(data.get(k) and data[k] == r['data'].get(k) for k in ('isbn_10', 'isbn_13')) and compatible(data, r['data']))), None)
            if same:
                classification = 'EXISTING_HOLDING'
            elif classification != 'EXISTING_HOLDING' and any(normalized_title(data['title']) == normalized_title(r['data']['title']) for r in pending):
                classification = 'POSSIBLE_DUPLICATE'
            result = {'row_number': item['row_number'], 'classification': classification, 'title': data['title'], 'subjects': data['subjects'],
                      'global_work_id': data.get('global_work_id'), 'isbn_10': data.get('isbn_10'), 'isbn_13': data.get('isbn_13'), 'data': data,
                      'reason': 'Possible existing item; this row creates a separate record unless a reliable identity matches.' if classification == 'POSSIBLE_DUPLICATE' else None,
                      'matched_library_book_id': match['library_book_id'] if match and classification == 'EXISTING_HOLDING' else None,
                      'matched_row_number': same['row_number'] if same else None}
            rows.append(result); pending.append(result)
        except (HTTPException, ValidationError, ValueError) as error:
            rows.append({'row_number': item['row_number'], 'classification': 'INVALID', 'reason': safe_reason(error), 'title': item['request'].title})
    require_readable(library_id)
    if len(json.dumps(rows, default=str).encode()) > 6 * 1024 * 1024:
        raise HTTPException(413, 'Normalized preview is too large. Split this inventory file.')
    import_id = str(uuid4()); counts = summary(rows)
    job = {'import_id': import_id, 'library_id': library_id, 'actor_user_id': actor_id(current_user), 'filename': name,
           'file_hash': hashlib.sha256(content).hexdigest(), 'status': 'PREVIEW', **counts, 'created_records': 0,
           'created_at': now(), 'started_at': None, 'completed_at': None, 'expires_at': now() + timedelta(hours=1)}
    repo.insert('inventory_previews', {**job, 'rows': rows})
    repo.insert('inventory_imports', job)
    return {**{k: v for k, v in job.items() if k != 'file_hash'}, 'rows': [{k: v for k, v in row.items() if k != 'data'} for row in rows]}


def apply(library_id, import_id, request, current_user):
    repo = TenantInventory(library_id)
    with library_mutation(library_id, wait_seconds=30):
        library, _ = require_library_staff(library_id, current_user, locked=True)
        from backend.services.library_lending_journal import state
        if state(library_id) == 'CANONICAL' and request.existing_policy == 'increase_quantity':
            raise HTTPException(409, 'Use canonical Manage Stock to increase existing holdings. Import new records with Skip existing holding.')
        job = repo.one('inventory_imports', {'import_id': import_id, 'actor_user_id': actor_id(current_user)})
        if not job:
            raise HTTPException(404, 'Import preview not found for this account and library.')
        if job['status'] in ('COMPLETED', 'PARTIAL'):
            return job['result']
        preview_row = repo.one('inventory_previews', {'import_id': import_id, 'actor_user_id': actor_id(current_user), 'file_hash': job['file_hash']})
        expiry = job['expires_at'].replace(tzinfo=timezone.utc) if job['expires_at'].tzinfo is None else job['expires_at']
        if not preview_row or expiry <= now():
            raise HTTPException(410, 'Inventory preview expired. Upload the file again.')
        repo.update('inventory_imports', {'import_id': import_id}, {'status': 'IMPORTING', 'started_at': now(), 'existing_policy': request.existing_policy})
        audit(library_id, 'INVENTORY_IMPORT_STARTED', actor_id(current_user), metadata={'import_id': import_id, 'counts': {'rows_total': job['rows_total']}})
        results = []; created = increased = skipped = invalid = 0
        candidates = isbn_candidates([row['data'].get(k) for row in preview_row['rows'] if 'data' in row for k in ('isbn_10', 'isbn_13')])
        for item in preview_row['rows']:
            if item['classification'] == 'INVALID':
                invalid += 1; results.append({**item, 'outcome': 'INVALID'}); continue
            try:
                data = item['data']
                value, added = insert_locked(repo, library, data, actor_id(current_user), source='DATASET_IMPORT',
                                            source_key=f'import:{import_id}:{item["row_number"]}', candidates=candidates, allow_existing=True)
                if added:
                    created += 1; outcome = 'CREATED'
                elif request.existing_policy == 'increase_quantity':
                    holding = value['holding']
                    if data.get('barcode') and repo.one('copies', {'barcode': data['barcode']}):
                        raise HTTPException(409, 'Barcode already belongs to a registered copy.')
                    total = holding['total_copies'] + data['copies']
                    available = holding['available_copies'] + (data['available_copies'] if data['available_copies'] is not None else data['copies'])
                    if total > 1000000:
                        raise HTTPException(422, 'Holding exceeds the supported quantity bound.')
                    repo.update('holdings', {'holding_id': holding['holding_id']}, {'total_copies': total, 'available_copies': available, 'updated_at': now(), 'updated_by_user_id': actor_id(current_user)})
                    if data.get('barcode'):
                        add_copy_locked(repo, value, CopyCreate(barcode=data['barcode']), actor_id(current_user))
                    audit(library_id, 'HOLDING_UPDATED', actor_id(current_user), metadata={'library_book_id': value['library_book_id'], 'import_id': import_id, 'counts': {'total': total, 'available_snapshot': available}})
                    increased += 1; outcome = 'QUANTITY_INCREASED'
                else:
                    skipped += 1; outcome = 'SKIPPED_EXISTING'
                results.append({'row_number': item['row_number'], 'classification': item['classification'], 'outcome': outcome,
                                'library_book_id': value['library_book_id'], 'accession_number': value['accession_number'], 'title': value['title']})
            except (HTTPException, ValidationError, ValueError) as error:
                invalid += 1; results.append({'row_number': item['row_number'], 'classification': 'INVALID', 'outcome': 'INVALID', 'reason': safe_reason(error), 'title': item['title']})
        status = 'PARTIAL' if invalid else 'COMPLETED'
        result = {'import_id': import_id, 'library_id': library_id, 'status': status, 'rows_total': len(results), 'created_records': created,
                  'increased_holdings': increased, 'skipped_existing': skipped, 'invalid_rows': invalid, 'possible_duplicates': job['possible_duplicates'], 'rows': results}
        audit(library_id, 'INVENTORY_IMPORT_PARTIAL' if invalid else 'INVENTORY_IMPORT_COMPLETED', actor_id(current_user),
              metadata={'import_id': import_id, 'counts': {'created': created, 'increased': increased, 'skipped': skipped, 'invalid': invalid}})
        repo.update('inventory_imports', {'import_id': import_id}, {'status': status, 'completed_at': now(), 'created_records': created, 'result': result})
    return result


def job_status(library_id, import_id, current_user):
    require_library_staff(library_id, current_user)
    row = TenantInventory(library_id).one('inventory_imports', {'import_id': import_id, 'actor_user_id': actor_id(current_user)})
    if not row:
        raise HTTPException(404, 'Import not found.')
    require_readable(library_id)
    return {k: v for k, v in row.items() if k not in ('_id', 'file_hash')}
